from flask import Flask, request, jsonify
from flask_cors import CORS
import pdfplumber
import docx
import os
import re
import threading
import unicodedata
from pathlib import Path

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": ["http://localhost:8000", "http://127.0.0.1:8000"]}})
RESUME_DIRECTORY = (Path(__file__).resolve().parent.parent / 'uploads' / 'resumes').resolve()
SENTENCE_MODEL_NAME = os.environ.get(
    'HIREBASE_EMBEDDING_MODEL',
    'sentence-transformers/all-MiniLM-L6-v2'
)
_sentence_model = None
_sentence_model_loaded = False
_sentence_model_lock = threading.Lock()


STOP_WORDS = {
    'a', 'an', 'and', 'are', 'as', 'at', 'be', 'by', 'can', 'for', 'from',
    'has', 'have', 'in', 'is', 'it', 'of', 'on', 'or', 'our', 'role', 'that',
    'the', 'their', 'this', 'to', 'was', 'we', 'will', 'with', 'work', 'you',
    'your', 'job', 'required', 'requirements', 'experience', 'candidate',
    'ability', 'skills', 'skill', 'years', 'year', 'strong', 'good', 'must',
    'knowledge', 'understanding', 'including', 'such', 'using', 'etc', 'need',
    'needs', 'seeking', 'looking', 'responsibilities', 'preferred', 'join',
    'team', 'teams'
}

SKILL_ALIASES = {
    'css': ('css', 'css3'),
    'html': ('html', 'html5'),
    'javascript': ('javascript', 'js'),
    'typescript': ('typescript', 'ts'),
    'node.js': ('node.js', 'nodejs', 'node js'),
    'react': ('react', 'react.js', 'reactjs'),
    'postgresql': ('postgresql', 'postgres', 'postgres sql'),
    'mysql': ('mysql', 'my sql'),
    'spring boot': ('spring boot', 'springboot'),
    'rest api': ('rest api', 'restful api'),
    'machine learning': ('machine learning', 'ml'),
    'csharp': ('csharp', 'c sharp', 'c#'),
    'cplusplus': ('cplusplus', 'c plus plus', 'c++'),
    'kubernetes': ('kubernetes', 'k8s'),
    'aws': ('aws', 'amazon web services'),
    'gcp': ('gcp', 'google cloud', 'google cloud platform'),
    'azure': ('azure', 'microsoft azure'),
    'dotnet': ('dotnet', '.net'),
    'sql': ('sql', 'structured query language'),
    'python': ('python',),
    'java': ('java',),
    'flask': ('flask',),
    'django': ('django',),
    'spring': ('spring',),
    'mongodb': ('mongodb', 'mongo db'),
    'redis': ('redis',),
    'docker': ('docker',),
    'git': ('git', 'github'),
    'linux': ('linux',),
    'data analysis': ('data analysis', 'data analytics'),
    'tensorflow': ('tensorflow',),
    'pytorch': ('pytorch',),
    'pandas': ('pandas',),
    'scikit-learn': ('scikit-learn', 'sklearn'),
    'excel': ('excel', 'microsoft excel'),
    'communication': ('communication', 'written communication'),
    'project management': ('project management',),
    'agile': ('agile', 'scrum'),
}

SECTION_PATTERNS = {
    'summary': r'\b(summary|professional summary|profile|objective)\b',
    'experience': r'\b(experience|employment|work history|professional history)\b',
    'education': r'\b(education|academic background|qualifications)\b',
    'skills': r'\b(skills|technical skills|core competencies)\b',
    'projects': r'\b(projects|personal projects|selected projects)\b',
    'certifications': r'\b(certifications?|licenses?)\b',
}

ACTION_VERBS = {
    'achieved', 'automated', 'built', 'collaborated', 'created', 'delivered',
    'designed', 'developed', 'improved', 'implemented', 'increased', 'launched',
    'led', 'managed', 'optimized', 'reduced', 'resolved', 'shipped', 'streamlined'
}


def normalize_text(value):
    """Normalize text for comparisons while preserving technical tokens."""
    value = unicodedata.normalize('NFKC', str(value or '')).lower()
    value = value.replace('c++', 'cplusplus').replace('c#', 'csharp')
    value = value.replace('.net', ' dotnet ')
    value = re.sub(r'(?<![a-z0-9])\.|\.(?![a-z0-9])', ' ', value)
    value = re.sub(r'[^a-z0-9+#.]+', ' ', value)
    return re.sub(r'\s+', ' ', value).strip()


def canonicalize_analysis_text(value):
    """Unify common versioned web skills while preserving other terms."""
    value = normalize_text(value)
    value = re.sub(r'\bcss\s*3\b', 'css', value)
    value = re.sub(r'\bhtml\s*5\b', 'html', value)
    return value


def canonical_skill(value):
    """Return a stable comparison key for equivalent skill spellings."""
    normalized = normalize_text(value)
    aliases = {
        'css3': 'css', 'html5': 'html',
        'js': 'javascript', 'nodejs': 'node.js', 'node js': 'node.js',
        'ts': 'typescript', 'reactjs': 'react', 'react.js': 'react',
        'postgres': 'postgresql', 'postgres sql': 'postgresql',
        'c sharp': 'csharp', 'c plus plus': 'cplusplus',
        'k8s': 'kubernetes', 'amazon web services': 'aws',
        'google cloud platform': 'gcp', 'google cloud': 'gcp',
        'microsoft azure': 'azure'
    }
    aliases.update({
        normalize_text(alias): canonical
        for canonical, spellings in SKILL_ALIASES.items()
        for alias in spellings
    })
    return aliases.get(normalized, normalized)


def skill_list(value):
    """Accept comma, semicolon, pipe, or newline separated skill lists."""
    if isinstance(value, list):
        parts = value
    else:
        parts = re.split(r'[,;|\n\r]+', str(value or ''))
    unique = {}
    for part in parts:
        display = re.sub(r'\s+', ' ', str(part)).strip()
        normalized = normalize_text(display)
        if normalized:
            normalized = canonical_skill(normalized)
            unique.setdefault(normalized, display)
    return unique


def keyword_list(job_text):
    normalized = normalize_text(job_text)
    tokens = re.findall(r'[a-z0-9]+(?:[+#.]?[a-z0-9]+)*', normalized)
    return list(dict.fromkeys(
        token for token in tokens
        if len(token) > 1 and token not in STOP_WORDS and not token.isdigit()
    ))


def contains_skill(text, skill):
    """Match a skill with word boundaries and known spelling aliases."""
    normalized = canonicalize_analysis_text(text)
    canonical = canonical_skill(skill)
    spellings = SKILL_ALIASES.get(canonical, (skill, canonical))
    for spelling in spellings:
        normalized_spelling = canonicalize_analysis_text(spelling)
        if normalized_spelling and re.search(
            rf'(?<![a-z0-9]){re.escape(normalized_spelling)}(?![a-z0-9])',
            normalized
        ):
            return True
    return False


def extract_known_skills(text):
    """Find supported skills from job/resume text using a small alias ontology."""
    return [
        skill for skill in SKILL_ALIASES
        if contains_skill(text, skill)
    ]


def tfidf_similarity(resume_text, job_text):
    """Combine word and character TF-IDF similarity for robust text matching."""
    word_similarity = 0.0
    char_similarity = 0.0
    try:
        word_vectorizer = TfidfVectorizer(
            ngram_range=(1, 2),
            max_features=1200,
            sublinear_tf=True,
            token_pattern=r'(?u)\b[a-zA-Z][a-zA-Z0-9+#.]*\b'
        )
        word_vectors = word_vectorizer.fit_transform([resume_text, job_text])
        word_similarity = float(cosine_similarity(word_vectors[0:1], word_vectors[1:2])[0][0])
    except ValueError:
        pass
    try:
        char_vectorizer = TfidfVectorizer(
            analyzer='char_wb', ngram_range=(3, 5), max_features=5000,
            sublinear_tf=True
        )
        char_vectors = char_vectorizer.fit_transform([resume_text, job_text])
        char_similarity = float(cosine_similarity(char_vectors[0:1], char_vectors[1:2])[0][0])
    except ValueError:
        pass
    return (word_similarity * 0.8 + char_similarity * 0.2) * 100


def split_evidence_sentences(text, limit=80):
    """Keep readable source excerpts for explainable and semantic comparisons."""
    sentences = []
    for line in str(text or '').splitlines():
        for sentence in re.split(r'(?<=[.!?])\s+', line):
            excerpt = re.sub(r'^\s*(?:[-*•]+\s*)+', '', sentence).strip()
            if len(excerpt) >= 16 and excerpt not in sentences:
                sentences.append(excerpt[:500])
            if len(sentences) >= limit:
                return sentences
    return sentences


def get_sentence_model():
    """Load a compact sentence-embedding model once, on the first analysis request."""
    global _sentence_model, _sentence_model_loaded
    if _sentence_model_loaded:
        return _sentence_model

    with _sentence_model_lock:
        if _sentence_model_loaded:
            return _sentence_model
        _sentence_model_loaded = True
        try:
            from sentence_transformers import SentenceTransformer
            _sentence_model = SentenceTransformer(SENTENCE_MODEL_NAME, device='cpu')
        except Exception as error:
            app.logger.warning(
                'Sentence embeddings unavailable (%s). Falling back to TF-IDF for this run.',
                error
            )
            _sentence_model = None
    return _sentence_model


def semantic_requirement_analysis(resume_text, job_text):
    """Compare job sentences with resume sentences using local sentence embeddings."""
    job_sentences = split_evidence_sentences(job_text, limit=40)
    resume_sentences = split_evidence_sentences(resume_text, limit=80)
    model = get_sentence_model()
    empty_result = {
        'available': False,
        'model': SENTENCE_MODEL_NAME,
        'score': None,
        'matches': [],
        'note': 'Sentence embeddings are unavailable; exact skill evidence and TF-IDF are still shown.'
    }
    if model is None or not job_sentences or not resume_sentences:
        return empty_result

    try:
        embeddings = model.encode(
            job_sentences + resume_sentences,
            convert_to_numpy=True,
            normalize_embeddings=True,
            show_progress_bar=False
        )
        job_vectors = embeddings[:len(job_sentences)]
        resume_vectors = embeddings[len(job_sentences):]
        similarities = cosine_similarity(job_vectors, resume_vectors)
        matches = []
        best_scores = []
        for index, job_sentence in enumerate(job_sentences):
            resume_index = int(similarities[index].argmax())
            similarity = max(0.0, min(1.0, float(similarities[index][resume_index])))
            best_scores.append(similarity)
            status = 'matched' if similarity >= 0.48 else 'related' if similarity >= 0.30 else 'missing'
            matches.append({
                'job_requirement': job_sentence,
                'resume_evidence': resume_sentences[resume_index] if status != 'missing' else None,
                'similarity': round(similarity * 100, 1),
                'status': status,
                'explanation': (
                    'Strongly related wording found in the resume.' if status == 'matched'
                    else 'Possibly related wording found; verify that it supports this requirement.' if status == 'related'
                    else 'No close sentence-level evidence was found in the extracted resume text.'
                )
            })

        matches.sort(key=lambda item: item['similarity'], reverse=True)
        score = (sum(best_scores) / len(best_scores) * 100) if best_scores else 0.0
        return {
            'available': True,
            'model': SENTENCE_MODEL_NAME,
            'score': round(score, 1),
            'matches': matches[:12],
            'note': 'Sentence-level semantic similarity finds related wording. Review each excerpt; it does not prove a skill or qualification.'
        }
    except Exception as error:
        app.logger.warning('Sentence embedding comparison failed: %s', error)
        empty_result['note'] = 'Sentence comparison failed; exact skill evidence and TF-IDF are still shown.'
        return empty_result


def exact_skill_evidence(resume_text, required_items):
    """Return the resume sentences that explicitly support each required skill."""
    resume_sentences = split_evidence_sentences(resume_text)
    evidence = []
    for key, label in required_items:
        excerpts = [sentence for sentence in resume_sentences if contains_skill(sentence, key)][:3]
        evidence.append({
            'requirement': label,
            'status': 'matched' if excerpts else 'missing',
            'resume_evidence': excerpts,
            'explanation': (
                'The resume explicitly mentions this skill (including recognized aliases).' if excerpts
                else 'No explicit mention was found. Related wording is shown separately and does not count as an exact skill match.'
            )
        })
    return evidence


def make_suggestion(priority, title, detail):
    return {'priority': priority, 'title': title, 'detail': detail}


# Week 3 - Resume text extraction
@app.route('/extract-resume-text', methods=['POST'])
def extract_resume_text():
    try:
        data = request.get_json(silent=True) or {}
        file_path = data.get('file_path')

        if not file_path:
            return jsonify({'error': 'No file path provided'}), 400

        requested_path = Path(file_path).resolve()
        if requested_path.parent != RESUME_DIRECTORY:
            return jsonify({'error': 'Resume path is outside the resume storage directory.'}), 403
        if not requested_path.is_file():
            return jsonify({'error': 'File not found'}), 404

        file_extension = requested_path.suffix.lower()

        extracted_text = ""

        if file_extension == '.pdf':
            with pdfplumber.open(requested_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        extracted_text += text + "\n"

        elif file_extension == '.docx':
            doc = docx.Document(requested_path)
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    extracted_text += ' | '.join(cell.text for cell in row.cells) + "\n"

        else:
            return jsonify({
                'error': 'Unsupported file format. Only PDF and DOCX are supported.'
            }), 400

        extracted_text = extracted_text.strip()
        if not extracted_text:
            return jsonify({'error': 'No selectable text was found. For a scanned PDF, use a text-based PDF or DOCX.'}), 422
        return jsonify({'text': extracted_text})

    except Exception as e:
        return jsonify({'error': str(e)}), 500


# Week 4 - Job and candidate matching
@app.route('/match-score', methods=['POST'])
def match_score():
    try:
        data = request.get_json(silent=True) or {}

        resume_text = str(data.get('resume_text') or '').strip()
        job_text = str(data.get('job_text') or '').strip()

        if not resume_text or not job_text:
            return jsonify({'error': 'Add resume text and job details to compare wording.'}), 400

        resume_text = canonicalize_analysis_text(resume_text)
        job_text = canonicalize_analysis_text(job_text)
        tfidf_score = tfidf_similarity(resume_text, job_text)
        semantic = semantic_requirement_analysis(resume_text, job_text)
        score = round(
            tfidf_score * .45 + semantic['score'] * .55
            if semantic['available'] else tfidf_score,
            2
        )

        return jsonify({
            'score': score,
            'tfidf_score': round(tfidf_score, 2),
            'semantic_similarity': semantic['score'],
            'semantic_available': semantic['available'],
            'semantic_matches': semantic['matches'],
            'score_type': 'resume_job_text_similarity',
            'score_note': 'Combines TF-IDF overlap with sentence embeddings when available. It is a relevance hint, not a hiring prediction.'
        })

    except Exception as e:
        return jsonify({'error': str(e)}), 500


# Week 4 - ATS resume score
@app.route('/ats-score', methods=['POST'])
def ats_score():
    try:
        data = request.get_json(silent=True) or {}
        resume_text = str(data.get('resume_text') or '').strip()
        job_text = str(data.get('job_text') or '').strip()
        if not resume_text or not job_text:
            return jsonify({'error': 'Add resume text and job details to calculate an estimate.'}), 400

        normalized_resume = canonicalize_analysis_text(resume_text)
        normalized_job = canonicalize_analysis_text(job_text)
        explicit_skills = skill_list(data.get('required_skills') or [])
        preferred_skills = skill_list(data.get('preferred_skills') or [])
        required_skills = explicit_skills or {
            skill: skill.replace('.', ' ').title()
            for skill in extract_known_skills(job_text)
        }
        required_items = list(required_skills.items())
        preferred_items = list(preferred_skills.items())
        matched_skills = [label for key, label in required_items if contains_skill(resume_text, key)]
        missing_skills = [label for key, label in required_items if not contains_skill(resume_text, key)]
        requirements_evidence = exact_skill_evidence(resume_text, required_items)
        preferred_evidence = exact_skill_evidence(resume_text, preferred_items)
        matched_preferred_skills = [label for key, label in preferred_items if contains_skill(resume_text, key)]
        missing_preferred_skills = [label for key, label in preferred_items if not contains_skill(resume_text, key)]
        skill_score = len(matched_skills) / len(required_items) * 100 if required_items else 0.0

        keywords = keyword_list(job_text)
        keyword_groups = {}
        for word in [*[key for key, _ in required_items], *keywords]:
            keyword_groups.setdefault(canonical_skill(word), word)
            if len(keyword_groups) >= 40:
                break
        matched_keywords = [label for key, label in keyword_groups.items() if contains_skill(resume_text, key)]
        missing_keywords = [label for key, label in keyword_groups.items() if not contains_skill(resume_text, key)]
        keyword_score = len(matched_keywords) / len(keyword_groups) * 100 if keyword_groups else 0.0

        sections_found = [
            name for name, pattern in SECTION_PATTERNS.items()
            if re.search(pattern, resume_text, re.IGNORECASE)
        ]
        sections_missing = [name for name in SECTION_PATTERNS if name not in sections_found]
        section_score = len(sections_found) / len(SECTION_PATTERNS) * 100
        text_similarity = tfidf_similarity(normalized_resume, normalized_job)
        semantic = semantic_requirement_analysis(resume_text, job_text)
        word_count = len(re.findall(r'\b\w+\b', resume_text))
        quantified_lines = sum(
            bool(re.search(
                r'\b\d+(?:[.,]\d+)?\s*(?:(?:percent|x|users?|customers?|clients?|projects?|hours?|days?|weeks?|months?|years?|k|m)\b|%)',
                line, re.IGNORECASE
            ))
            for line in resume_text.splitlines() if line.strip()
        )
        action_verb_count = sum(
            len(re.findall(rf'\b{re.escape(verb)}\b', resume_text, re.IGNORECASE))
            for verb in ACTION_VERBS
        )
        email_found = bool(re.search(r'\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b', resume_text, re.IGNORECASE))
        phone_found = bool(re.search(r'(?<!\w)\+?\d[\d ()-]{7,}\d(?!\w)', resume_text))

        dimensions = {
            'skill_coverage': round(skill_score, 1) if required_items else None,
            'keyword_coverage': round(keyword_score, 1),
            'text_similarity': round(text_similarity, 1),
            'semantic_similarity': semantic['score'],
            'resume_structure': round(section_score, 1),
        }
        if semantic['available'] and required_items:
            weights = {'skill_coverage': 40, 'keyword_coverage': 25, 'text_similarity': 15, 'semantic_similarity': 10, 'resume_structure': 10}
            final_score = skill_score * .40 + keyword_score * .25 + text_similarity * .15 + semantic['score'] * .10 + section_score * .10
        elif semantic['available']:
            weights = {'skill_coverage': 0, 'keyword_coverage': 45, 'text_similarity': 25, 'semantic_similarity': 20, 'resume_structure': 10}
            final_score = keyword_score * .45 + text_similarity * .25 + semantic['score'] * .20 + section_score * .10
        elif required_items:
            weights = {'skill_coverage': 45, 'keyword_coverage': 25, 'text_similarity': 20, 'semantic_similarity': 0, 'resume_structure': 10}
            final_score = skill_score * .45 + keyword_score * .25 + text_similarity * .20 + section_score * .10
        else:
            weights = {'skill_coverage': 0, 'keyword_coverage': 50, 'text_similarity': 30, 'semantic_similarity': 0, 'resume_structure': 20}
            final_score = keyword_score * .50 + text_similarity * .30 + section_score * .20

        suggestions = []
        if missing_skills:
            suggestions.append(make_suggestion(
                'high', 'Address required skills',
                'If you have this experience, show it with a specific project or work example: ' + ', '.join(missing_skills[:8]) + '.'
            ))
        elif required_items:
            suggestions.append(make_suggestion(
                'positive', 'Required skills detected',
                'Your resume contains all listed skills detected for this role. Keep the supporting examples specific.'
            ))
        if missing_preferred_skills:
            suggestions.append(make_suggestion(
                'low', 'Consider preferred skills',
                'These are optional qualifications. If you have relevant experience, make it visible: '
                + ', '.join(missing_preferred_skills[:8]) + '.'
            ))
        if missing_keywords:
            suggestions.append(make_suggestion(
                'medium', 'Improve role keyword coverage',
                'Where accurate, include relevant job terms in your skills or experience: ' + ', '.join(missing_keywords[:8]) + '.'
            ))
        if sections_missing:
            suggestions.append(make_suggestion(
                'medium', 'Make sections easy to scan',
                'Add clear headings for relevant sections: ' + ', '.join(sections_missing) + '.'
            ))
        if quantified_lines == 0:
            suggestions.append(make_suggestion(
                'medium', 'Show measurable outcomes',
                'Add a truthful number, scale, time saved, users served, or percentage to one or two achievement bullets.'
            ))
        if action_verb_count < 3:
            suggestions.append(make_suggestion(
                'low', 'Strengthen experience bullets',
                'Start relevant bullets with specific verbs such as built, improved, automated, or delivered.'
            ))
        if word_count < 120:
            suggestions.append(make_suggestion(
                'medium', 'Add more resume detail',
                'The extracted resume is very short. Include relevant projects, tools, responsibilities, and outcomes if you have them.'
            ))
        elif word_count > 1400:
            suggestions.append(make_suggestion(
                'low', 'Tighten resume length',
                'Prioritize recent relevant experience and remove repeated details.'
            ))
        if text_similarity < 12:
            suggestions.append(make_suggestion(
                'medium', 'Connect experience to the role',
                'The wording overlap is low. Add concise, truthful examples that explain how your work relates to this role.'
            ))
        if not email_found or not phone_found:
            missing_contact = 'an email address' if not email_found and phone_found else 'a phone number' if email_found and not phone_found else 'an email address and phone number'
            suggestions.append(make_suggestion(
                'low', 'Check your contact details',
                'The extracted resume text does not show ' + missing_contact + '. Confirm your preferred contact details are present and readable.'
            ))
        if not suggestions:
            suggestions.append(make_suggestion(
                'positive', 'Good alignment signals',
                'The resume has strong detected overlap. Review the original job post and keep every claim accurate.'
            ))

        score = round(max(0.0, min(100.0, final_score)), 1)
        score_label = 'Strong alignment' if score >= 75 else 'Good starting point' if score >= 55 else 'More tailoring recommended'
        return jsonify({
            'score': score,
            'score_label': score_label,
            'score_type': 'resume_alignment_estimate',
            'score_note': 'This local estimate combines exact skill and keyword checks, TF-IDF, sentence embeddings when available, and resume structure signals. It is not a hiring prediction or an employer ATS result.',
            'model': 'Sentence Transformers embeddings + scikit-learn TF-IDF + curated skill aliases' if semantic['available'] else 'scikit-learn TF-IDF + curated skill aliases (sentence embeddings unavailable)',
            'dimensions': dimensions,
            'weights': weights,
            'skill_score': round(skill_score, 1) if required_items else None,
            'keyword_score': round(keyword_score, 1),
            'section_score': round(section_score, 1),
            'text_similarity': round(text_similarity, 1),
            'required_skills': [label for _, label in required_items],
            'matched_skills': matched_skills,
            'missing_skills': missing_skills,
            'requirements_evidence': requirements_evidence,
            'preferred_skills': [label for _, label in preferred_items],
            'matched_preferred_skills': matched_preferred_skills,
            'missing_preferred_skills': missing_preferred_skills,
            'preferred_evidence': preferred_evidence,
            'semantic_analysis': semantic,
            'matched_keywords': matched_keywords,
            'missing_keywords': missing_keywords,
            'sections_found': sections_found,
            'sections_missing': sections_missing,
            'signals': {
                'resume_word_count': word_count,
                'sections_found': len(sections_found),
                'sections_total': len(SECTION_PATTERNS),
                'quantified_lines': quantified_lines,
                'action_verb_uses': action_verb_count,
                'email_found': email_found,
                'phone_found': phone_found,
                'job_terms_compared': len(keyword_groups),
            },
            'suggestions': suggestions,
        })

    except Exception as e:
        return jsonify({'error': str(e)}), 500


# Week 4 - Skill gap analysis
@app.route('/skill-gap', methods=['POST'])
def skill_gap():
    try:
        data = request.get_json(silent=True) or {}
        candidate_skills = skill_list(data.get('candidate_skills', ''))
        required_skills = skill_list(data.get('required_skills', ''))
        preferred_skills = skill_list(data.get('preferred_skills', ''))

        if not required_skills and not preferred_skills:
            return jsonify({'error': 'This job has no required or preferred skills to compare yet.'}), 400

        matched_keys = [key for key in required_skills if key in candidate_skills]
        missing_keys = [key for key in required_skills if key not in candidate_skills]
        matched_skills = [required_skills[key] for key in matched_keys]
        missing_skills = [required_skills[key] for key in missing_keys]
        matched_preferred_keys = [key for key in preferred_skills if key in candidate_skills]
        missing_preferred_keys = [key for key in preferred_skills if key not in candidate_skills]
        matched_preferred = [preferred_skills[key] for key in matched_preferred_keys]
        missing_preferred = [preferred_skills[key] for key in missing_preferred_keys]
        coverage = round(len(matched_keys) / len(required_skills) * 100, 1) if required_skills else None
        preferred_coverage = (
            round(len(matched_preferred_keys) / len(preferred_skills) * 100, 1)
            if preferred_skills else None
        )

        learning_plan = [
            {
                'skill': required_skills[key],
                'category': 'required',
                'priority': 'high',
                'suggested_step': f'Build a small, verifiable project using {required_skills[key]} and add the outcome to your resume.'
            }
            for key in missing_keys
        ] + [
            {
                'skill': preferred_skills[key],
                'category': 'preferred',
                'priority': 'medium',
                'suggested_step': f'Explore {preferred_skills[key]} through a guided exercise or project if it fits your career goals.'
            }
            for key in missing_preferred_keys
        ]

        suggestion = (
            'All required skills are present in your profile.' if not missing_skills
            else 'Start with required skills. Add a skill to your profile only after you have experience you can support.'
        )

        return jsonify({
            'coverage': coverage,
            'matched_skills': matched_skills,
            'missing_skills': missing_skills,
            'preferred_coverage': preferred_coverage,
            'matched_preferred_skills': matched_preferred,
            'missing_preferred_skills': missing_preferred,
            'learning_plan': learning_plan,
            'suggestion': suggestion,
            'required_count': len(required_skills),
            'matched_count': len(matched_skills),
            'preferred_count': len(preferred_skills),
            'matched_preferred_count': len(matched_preferred)
        })

    except Exception as e:
        return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
    debug = os.environ.get('HIREBASE_ML_DEBUG', '').strip().lower() == 'true'
    app.run(host='localhost', port=5000, debug=debug)
